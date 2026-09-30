import io
from decimal import Decimal
from datetime import timedelta, date
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone
from django.core.exceptions import PermissionDenied

import openpyxl

from apps.accounts.models import User
from apps.organizations.models import Department, Location, Employee
from apps.inventory.models import Asset, AssetCategory, Brand, AssetVerification
from apps.assignments.models import AssetAssignment
from apps.transfers.models import AssetTransfer
from apps.borrowing.models import AssetBorrowing
from apps.maintenance.models import AssetMaintenance
from apps.supplies.models import Supply, SupplyCategory, SupplyTransaction
from apps.disposals.models import AssetDisposal
from apps.audit.models import AuditLog

from . import services, exporters


class Phase10ReportsBaseTestCase(TestCase):
    def setUp(self):
        self.client = Client()

        # Users
        self.admin = User.objects.create_user(
            username='admin_rep',
            password='password123',
            role=User.Role.ADMIN
        )
        self.dean = User.objects.create_user(
            username='dean_rep',
            password='password123',
            role=User.Role.DEAN
        )
        self.chair_ba = User.objects.create_user(
            username='chair_ba',
            password='password123',
            role=User.Role.DEPT_CHAIR
        )
        self.chair_it = User.objects.create_user(
            username='chair_it',
            password='password123',
            role=User.Role.DEPT_CHAIR
        )
        self.faculty = User.objects.create_user(
            username='faculty_rep',
            password='password123',
            role=User.Role.FACULTY
        )

        # Departments
        self.dept_ba = Department.objects.create(name='Business Administration', code='BA')
        self.dept_it = Department.objects.create(name='Information Technology', code='IT')

        # Locations
        self.loc_ba = Location.objects.create(name='BA Dean Office', building='Main', room_number='101', department=self.dept_ba)
        self.loc_it = Location.objects.create(name='IT Lab 1', building='IT Complex', room_number='201', department=self.dept_it)

        # Employees
        self.emp_chair_ba = Employee.objects.create(
            user=self.chair_ba,
            employee_id='EMP-CBA-001',
            first_name='BA',
            last_name='Chair',
            department=self.dept_ba,
            position='Department Chair'
        )
        self.emp_chair_it = Employee.objects.create(
            user=self.chair_it,
            employee_id='EMP-CIT-001',
            first_name='IT',
            last_name='Chair',
            department=self.dept_it,
            position='Department Chair'
        )
        self.emp_faculty = Employee.objects.create(
            user=self.faculty,
            employee_id='EMP-FAC-001',
            first_name='Juan',
            last_name='Dela Cruz',
            department=self.dept_ba,
            position='Instructor'
        )

        # Categories & Brands
        self.cat_comp = AssetCategory.objects.create(name='Computers', code='COMP')
        self.cat_furn = AssetCategory.objects.create(name='Furniture', code='FURN')
        self.brand_dell = Brand.objects.create(name='Dell')

        # Assets
        self.asset_active_ba = Asset.objects.create(
            asset_code='CBA-COMP-001',
            property_number='PROP-2026-001',
            item_name='Dell Latitude Laptop',
            category=self.cat_comp,
            brand=self.brand_dell,
            department=self.dept_ba,
            current_location=self.loc_ba,
            condition=Asset.Condition.GOOD,
            status=Asset.Status.AVAILABLE,
            acquisition_cost=Decimal('45000.00'),
            acquisition_date=date(2025, 1, 15)
        )
        self.asset_assigned_ba = Asset.objects.create(
            asset_code='CBA-COMP-002',
            property_number='PROP-2026-002',
            item_name='Dell OptiPlex Desktop',
            category=self.cat_comp,
            brand=self.brand_dell,
            department=self.dept_ba,
            current_location=self.loc_ba,
            condition=Asset.Condition.NEW,
            status=Asset.Status.ASSIGNED,
            acquisition_cost=Decimal('50000.00'),
            acquisition_date=date(2025, 2, 1)
        )
        self.asset_it = Asset.objects.create(
            asset_code='CIT-COMP-001',
            property_number='PROP-2026-003',
            item_name='Server Workstation',
            category=self.cat_comp,
            brand=self.brand_dell,
            department=self.dept_it,
            current_location=self.loc_it,
            condition=Asset.Condition.GOOD,
            status=Asset.Status.AVAILABLE,
            acquisition_cost=Decimal('80000.00'),
            acquisition_date=date(2025, 3, 10)
        )
        self.asset_disposed = Asset.objects.create(
            asset_code='CBA-FURN-001',
            property_number='PROP-2024-099',
            item_name='Broken Conference Table',
            category=self.cat_furn,
            department=self.dept_ba,
            current_location=self.loc_ba,
            condition=Asset.Condition.UNSERVICEABLE,
            status=Asset.Status.DISPOSED,
            acquisition_cost=Decimal('15000.00'),
            acquisition_date=date(2024, 6, 1)
        )

        # Assignment
        self.assignment_active = AssetAssignment.objects.create(
            asset=self.asset_assigned_ba,
            employee=self.emp_faculty,
            assigned_by=self.admin,
            assigned_date=timezone.now() - timedelta(days=30),
            condition_at_assignment=Asset.Condition.NEW,
            status=AssetAssignment.Status.ACTIVE
        )


class TestReportsSecurityAndScoping(Phase10ReportsBaseTestCase):
    """Verifies RBAC rules and unbreakable QuerySet scoping."""

    def test_unauthenticated_redirect(self):
        url = reverse('reports:reports_dashboard')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 302)
        self.assertIn('/accounts/login/', response.url)

    def test_faculty_blocked_with_403(self):
        """Faculty members must be blocked with HTTP 403 on all administrative reports."""
        self.client.force_login(self.faculty)

        urls_to_test = [
            reverse('reports:reports_dashboard'),
            reverse('reports:asset_inventory'),
            reverse('reports:asset_by_department'),
            reverse('reports:accountability_active'),
            reverse('reports:transfers'),
            reverse('reports:borrowing'),
            reverse('reports:maintenance'),
            reverse('reports:verification'),
            reverse('reports:supplies_stock'),
            reverse('reports:disposals'),
            reverse('reports:audit_log'),
            reverse('reports:executive_summary_xlsx'),
        ]
        for url in urls_to_test:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 403, f"Faculty accessed {url} without 403!")

    def test_dept_chair_strictly_scoped_at_queryset_level(self):
        """
        Department Chair MUST be scoped strictly to their own department.
        Attempting to query other departments (?department=IT) MUST NOT escape scope.
        """
        self.client.force_login(self.chair_ba)

        # Normal view
        url = reverse('reports:asset_inventory')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        items = list(response.context['assets'])
        self.assertIn(self.asset_active_ba, items)
        self.assertIn(self.asset_assigned_ba, items)
        self.assertNotIn(self.asset_it, items)

        # Attempted escape via GET parameter
        response_escape = self.client.get(url + f"?department={self.dept_it.pk}")
        self.assertEqual(response_escape.status_code, 200)
        items_escape = list(response_escape.context['assets'])
        # The IT asset must NOT be returned, despite ?department=IT
        self.assertNotIn(self.asset_it, items_escape)
        # Only BA assets remain in the scoped result
        for item in items_escape:
            self.assertEqual(item.department, self.dept_ba)

    def test_dept_chair_cannot_view_other_dept_employee_dossier(self):
        """Department Chair cannot view employee accountability outside their department."""
        self.client.force_login(self.chair_ba)
        # Attempt to access IT Chair dossier
        url = reverse('reports:accountability_employee', kwargs={'pk': self.emp_chair_it.pk})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

    def test_dept_chair_forbidden_from_audit_log(self):
        """Department Chair is strictly forbidden from system audit log report."""
        self.client.force_login(self.chair_ba)
        url = reverse('reports:audit_log')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

    def test_dean_has_college_wide_access_but_no_audit_log(self):
        """Dean can view all departmental records, but is blocked from system audit log."""
        self.client.force_login(self.dean)

        # Dean can view inventory across all departments
        url = reverse('reports:asset_inventory')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        items = list(response.context['assets'])
        self.assertIn(self.asset_active_ba, items)
        self.assertIn(self.asset_it, items)

        # Dean is blocked from audit log
        audit_url = reverse('reports:audit_log')
        audit_res = self.client.get(audit_url)
        self.assertEqual(audit_res.status_code, 403)

    def test_admin_has_full_access(self):
        """Admin can access all reports including the audit log."""
        self.client.force_login(self.admin)
        url = reverse('reports:audit_log')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)


class TestAssetInventoryReporting(Phase10ReportsBaseTestCase):
    """Verifies asset inventory active/disposed scope, date filtering, and summaries."""

    def test_active_scope_excludes_disposed(self):
        qs, summary = services.get_asset_inventory_report(self.admin, {'scope': 'active'})
        self.assertIn(self.asset_active_ba, qs)
        self.assertIn(self.asset_it, qs)
        self.assertNotIn(self.asset_disposed, qs)
        self.assertEqual(summary['disposed'], 0)

    def test_disposed_scope_only_returns_disposed(self):
        qs, summary = services.get_asset_inventory_report(self.admin, {'scope': 'disposed'})
        self.assertEqual(qs.count(), 1)
        self.assertEqual(qs.first(), self.asset_disposed)
        self.assertEqual(summary['disposed'], 1)

    def test_all_scope_returns_everything(self):
        qs, summary = services.get_asset_inventory_report(self.admin, {'scope': 'all'})
        self.assertEqual(qs.count(), 4)
        self.assertIn(self.asset_disposed, qs)
        self.assertEqual(summary['disposed'], 1)

    def test_date_range_filtering(self):
        qs, _ = services.get_asset_inventory_report(self.admin, {
            'date_from': '2025-01-01',
            'date_to': '2025-01-31'
        })
        self.assertEqual(qs.count(), 1)
        self.assertEqual(qs.first(), self.asset_active_ba)

    def test_department_summary_service(self):
        data = services.get_asset_by_department_summary(self.admin)
        ba_data = next((d for d in data if d['department'] == self.dept_ba), None)
        self.assertIsNotNone(ba_data)
        self.assertEqual(ba_data['counts']['total'], 3)
        self.assertEqual(ba_data['counts']['active'], 2)
        self.assertEqual(ba_data['counts']['disposed'], 1)


class TestFacultyAccountabilityReporting(Phase10ReportsBaseTestCase):
    """Verifies active vs historical assignment tracking."""

    def test_active_accountability_report(self):
        qs, summary = services.get_accountability_report(self.admin, {}, historical=False)
        self.assertEqual(qs.count(), 1)
        self.assertEqual(qs.first(), self.assignment_active)
        self.assertEqual(summary['total_assignments'], 1)
        self.assertEqual(summary['unique_employees'], 1)

    def test_historical_turnover_report(self):
        # Return the active assignment
        self.assignment_active.status = AssetAssignment.Status.RETURNED
        self.assignment_active.returned_date = timezone.now()
        self.assignment_active.condition_at_return = Asset.Condition.GOOD
        self.assignment_active.returned_by = self.admin
        self.assignment_active.save()

        # Active report should be empty
        active_qs, _ = services.get_accountability_report(self.admin, {}, historical=False)
        self.assertEqual(active_qs.count(), 0)

        # Historical report should contain it
        hist_qs, summary = services.get_accountability_report(self.admin, {}, historical=True)
        self.assertEqual(hist_qs.count(), 1)
        self.assertEqual(summary['total_records'], 1)


class TestConsumableSuppliesReporting(Phase10ReportsBaseTestCase):
    """Verifies transaction-derived stock and unit-isolated consumption."""

    def setUp(self):
        super().setUp()
        self.cat_off = SupplyCategory.objects.create(name='Office Supplies', code='OFF')
        self.supply_paper = Supply.objects.create(
            supply_code='SUP-001',
            item_name='Bond Paper A4',
            category=self.cat_off,
            unit='reams',
            reorder_level=10
        )
        self.supply_pen = Supply.objects.create(
            supply_code='SUP-002',
            item_name='Ballpoint Pen Black',
            category=self.cat_off,
            unit='pcs',
            reorder_level=20
        )

        # Stock In: 15 reams of paper, 50 pens
        SupplyTransaction.objects.create(
            supply=self.supply_paper,
            transaction_type=SupplyTransaction.TransactionType.STOCK_IN,
            quantity=15,
            processed_by=self.admin
        )
        SupplyTransaction.objects.create(
            supply=self.supply_pen,
            transaction_type=SupplyTransaction.TransactionType.STOCK_IN,
            quantity=50,
            processed_by=self.admin
        )

        # Stock Out: 8 reams of paper to BA Dept
        SupplyTransaction.objects.create(
            supply=self.supply_paper,
            transaction_type=SupplyTransaction.TransactionType.STOCK_OUT,
            quantity=8,
            department=self.dept_ba,
            employee=self.emp_faculty,
            processed_by=self.admin
        )

    def test_stock_derived_dynamically_from_transactions(self):
        """Stock must be calculated from transactions (15 - 8 = 7 reams)."""
        supplies, summary = services.get_supply_stock_report(self.admin, {})
        paper = next(s for s in supplies if s.pk == self.supply_paper.pk)
        self.assertEqual(paper.calculated_stock, 7)
        self.assertEqual(paper.stock_status, 'LOW_STOCK')  # 7 <= 10 reorder level
        self.assertEqual(paper.shortfall, 3)

        pen = next(s for s in supplies if s.pk == self.supply_pen.pk)
        self.assertEqual(pen.calculated_stock, 50)
        self.assertEqual(pen.stock_status, 'IN_STOCK')

    def test_department_usage_isolated_by_unit(self):
        """Department usage must report 8 reams without mixing unit types."""
        usage = services.get_supply_department_usage_report(self.admin, {'department': str(self.dept_ba.pk)})
        self.assertEqual(len(usage), 1)
        self.assertEqual(usage[0]['department__name'], self.dept_ba.name)
        self.assertEqual(usage[0]['supply__unit'], 'reams')
        self.assertEqual(usage[0]['total_issued'], 8)


class TestBorrowingAndMaintenanceReporting(Phase10ReportsBaseTestCase):
    """Verifies overdue borrowing detection and maintenance cost calculations."""

    def test_borrowing_realtime_overdue_detection(self):
        # Create an active loan whose requested return date is in the past
        past_date = timezone.now() - timedelta(days=2)
        loan = AssetBorrowing.objects.create(
            asset=self.asset_active_ba,
            borrower=self.emp_faculty,
            borrower_department=self.dept_ba,
            requested_start=timezone.now() - timedelta(days=5),
            requested_return=past_date,
            status=AssetBorrowing.Status.RELEASED,
            purpose='Lecture presentation'
        )

        # Overdue report should pick this up immediately
        qs, summary = services.get_borrowing_report(self.admin, {}, overdue_only=True)
        self.assertEqual(qs.count(), 1)
        self.assertEqual(qs.first().pk, loan.pk)
        self.assertTrue(qs.first().is_overdue)
        self.assertEqual(summary['overdue'], 1)

    def test_maintenance_cost_summary_aggregation(self):
        AssetMaintenance.objects.create(
            case_number='MNT-2026-001',
            asset=self.asset_active_ba,
            issue_description='Motherboard replacement',
            severity=AssetMaintenance.Severity.HIGH,
            status=AssetMaintenance.Status.COMPLETED,
            repair_cost=Decimal('3500.00'),
            service_provider='TechFix Bayawan',
            repair_completed_at=timezone.now(),
            final_condition=Asset.Condition.GOOD,
            reported_by=self.admin
        )
        AssetMaintenance.objects.create(
            case_number='MNT-2026-002',
            asset=self.asset_assigned_ba,
            issue_description='RAM upgrade',
            severity=AssetMaintenance.Severity.LOW,
            status=AssetMaintenance.Status.COMPLETED,
            repair_cost=Decimal('1500.00'),
            service_provider='In-House IT',
            repair_completed_at=timezone.now(),
            final_condition=Asset.Condition.GOOD,
            reported_by=self.admin
        )

        qs, summary, by_dept, top = services.get_maintenance_cost_summary(self.admin, {})
        self.assertEqual(summary['completed_repairs_count'], 2)
        self.assertEqual(summary['total_cost'], Decimal('5000.00'))
        self.assertEqual(summary['average_cost'], Decimal('2500.00'))
        self.assertEqual(len(by_dept), 1)
        self.assertEqual(by_dept[0]['total_cost'], Decimal('5000.00'))


class TestPhysicalVerificationReporting(Phase10ReportsBaseTestCase):
    """Verifies audit mismatch and never-verified tracking."""

    def test_mismatches_mode(self):
        AssetVerification.objects.create(
            asset=self.asset_active_ba,
            verified_by=self.admin,
            expected_department=self.dept_ba,
            observed_department=self.dept_it,  # Location mismatch!
            expected_location=self.loc_ba,
            observed_location=self.loc_it,
            expected_condition=Asset.Condition.GOOD,
            observed_condition=Asset.Condition.GOOD,
            result=AssetVerification.VerificationResult.LOCATION_MISMATCH
        )

        records, summary = services.get_verification_report(self.admin, {}, mode='mismatches')
        self.assertEqual(len(records), 1)
        self.assertEqual(summary['location_mismatch'], 1)

    def test_never_verified_mode(self):
        records, summary = services.get_verification_report(self.admin, {}, mode='never_verified')
        # All 3 active assets were created with no verifications
        self.assertEqual(len(records), 3)
        self.assertEqual(summary['never_verified'], 3)


class TestExportersAndFormulaSanitization(Phase10ReportsBaseTestCase):
    """Verifies formula injection neutralization, CSV export, and XLSX workbooks."""

    def test_formula_injection_sanitization(self):
        dangerous_values = [
            "=SUM(A1:A10)",
            "+cmd|' /C calc'!A0",
            "-2+3*[cmd.exe]",
            "@SUM(1,1)",
            "\tTabSeparated",
            "\rCarriageReturn",
        ]
        for val in dangerous_values:
            sanitized = exporters.sanitize_for_spreadsheet(val)
            self.assertTrue(
                sanitized.startswith("'"),
                f"Formula {val} was not escaped with single quote! Got: {sanitized}"
            )

        safe_values = [
            "Normal Asset Name",
            "PROP-2026-001",
            12345,
            Decimal('99.99'),
            date(2026, 1, 1),
            None,
        ]
        for val in safe_values:
            sanitized = exporters.sanitize_for_spreadsheet(val)
            self.assertEqual(sanitized, val)

    def test_csv_export_endpoint(self):
        self.client.force_login(self.admin)
        url = reverse('reports:asset_inventory') + "?export=csv"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn('attachment; filename="CBA_Asset_Inventory_', response['Content-Disposition'])

        content = response.content.decode('utf-8-sig')  # Decodes UTF-8 with BOM
        self.assertIn("Asset Inventory Report", content)
        self.assertIn("CBA-COMP-001", content)
        self.assertIn("Dell Latitude Laptop", content)

    def test_xlsx_export_endpoint(self):
        self.client.force_login(self.admin)
        url = reverse('reports:asset_inventory') + "?export=xlsx"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment; filename="CBA_Asset_Inventory_', response['Content-Disposition'])

        # Verify openpyxl can read the returned bytes
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("Inventory", wb.sheetnames)
        ws = wb["Inventory"]
        # Metadata check (Title is on row 3)
        self.assertIn("Asset Inventory Report", str(ws.cell(row=3, column=1).value))

    def test_executive_summary_multi_sheet_xlsx(self):
        self.client.force_login(self.admin)
        url = reverse('reports:executive_summary_xlsx')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        expected_sheets = [
            "Executive Summary", "Asset Inventory", "Faculty Accountability",
            "Equipment Borrowing", "Maintenance & Repairs", "Consumable Supplies", "Asset Disposals"
        ]
        for sheet in expected_sheets:
            self.assertIn(sheet, wb.sheetnames, f"Sheet {sheet} missing from executive workbook!")
