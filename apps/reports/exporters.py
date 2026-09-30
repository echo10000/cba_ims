import csv
import io
from decimal import Decimal
from datetime import date, datetime
from django.http import HttpResponse, StreamingHttpResponse
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


DANGEROUS_FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def sanitize_for_spreadsheet(value):
    """
    Prevents spreadsheet formula injection (CSV/Excel Formula Injection).
    If a string value starts with =, +, -, @, tab, or carriage return,
    or has leading whitespace followed by a dangerous prefix,
    prepends a single quote (') to force spreadsheet processors to treat it as plain text.
    """
    if isinstance(value, str) and value:
        # Check raw leading character (covers raw \t, \r, =, +, -, @)
        if value[0] in DANGEROUS_FORMULA_PREFIXES:
            return f"'{value}"
        # Check stripped value (covers '   =cmd')
        stripped = value.strip()
        if stripped and stripped[0] in ('=', '+', '-', '@'):
            return f"'{value}"
    return value


def export_to_csv(filename, title, metadata, headers, rows):
    """
    Generates a professional, UTF-8 CSV HttpResponse.
    Includes report metadata banner, column headers, and sanitized rows.
    """
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    # Add UTF-8 BOM so Excel on Windows recognizes UTF-8 correctly
    response.write('\ufeff')

    writer = csv.writer(response)

    # 1. Institutional Title & Metadata Header
    writer.writerow(["NEGROS ORIENTAL STATE UNIVERSITY - Bayawan-Santa Catalina Campus"])
    writer.writerow(["College of Business Administration - Asset & Inventory Management System"])
    writer.writerow([title])
    writer.writerow([])

    if metadata:
        for label, val in metadata.items():
            writer.writerow([f"{label}:", val])
        writer.writerow([])

    # 2. Table Column Headers
    writer.writerow(headers)

    # 3. Data Rows
    for row in rows:
        sanitized_row = [sanitize_for_spreadsheet(cell) for cell in row]
        writer.writerow(sanitized_row)

    return response


def export_to_xlsx(filename, title, metadata, headers, rows, sheet_name='Report', column_types=None):
    """
    Generates a styled Microsoft Excel (.xlsx) workbook using openpyxl.
    Features:
    - Institutional header block
    - Active filters / metadata block
    - Dark navy table headers (#1B365D) with bold white text
    - Auto-filtered columns & freeze panes
    - Auto-adjusted column widths
    - Formatted numbers, currencies (PHP), and dates
    - Formula injection protection on user-supplied strings
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel sheet title limit

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Styling Palettes
    navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1B365D")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="555555")
    meta_label_font = Font(name="Calibri", size=10, bold=True, color="333333")
    meta_val_font = Font(name="Calibri", size=10, color="555555")
    cell_font = Font(name="Calibri", size=10)

    thin_border_side = Side(border_style="thin", color="D3D3D3")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(
        left=Side(border_style="thin", color="0D1E36"),
        right=Side(border_style="thin", color="0D1E36"),
        top=Side(border_style="medium", color="0D1E36"),
        bottom=Side(border_style="medium", color="0D1E36")
    )

    current_row = 1

    # 1. Title Block
    ws.cell(row=current_row, column=1, value="NEGROS ORIENTAL STATE UNIVERSITY — Bayawan-Santa Catalina Campus").font = title_font
    current_row += 1
    ws.cell(row=current_row, column=1, value="College of Business Administration — Asset & Inventory Management System").font = subtitle_font
    current_row += 1
    ws.cell(row=current_row, column=1, value=title).font = Font(name="Calibri", size=12, bold=True, color="000000")
    current_row += 2

    # 2. Metadata Block
    if metadata:
        for label, val in metadata.items():
            ws.cell(row=current_row, column=1, value=f"{label}:").font = meta_label_font
            ws.cell(row=current_row, column=2, value=str(val)).font = meta_val_font
            current_row += 1
        current_row += 1

    # 3. Table Headers
    header_row_idx = current_row
    for col_idx, header in enumerate(headers, start=1):
        c = ws.cell(row=header_row_idx, column=col_idx, value=header)
        c.fill = navy_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = header_border

    ws.row_dimensions[header_row_idx].height = 28

    # 4. Data Rows
    current_row += 1
    start_data_row = current_row

    for row in rows:
        ws.row_dimensions[current_row].height = 20
        for col_idx, cell_value in enumerate(row, start=1):
            c = ws.cell(row=current_row, column=col_idx)
            c.font = cell_font
            c.border = cell_border

            # Format detection
            col_type = None
            if column_types and (col_idx - 1) < len(column_types):
                col_type = column_types[col_idx - 1]

            if cell_value is None:
                c.value = "—"
                c.alignment = Alignment(horizontal="center")
            elif isinstance(cell_value, (Decimal, float)):
                c.value = float(cell_value)
                if col_type == 'currency':
                    c.number_format = '₱#,##0.00'
                    c.alignment = Alignment(horizontal="right")
                elif col_type == 'integer':
                    c.number_format = '#,##0'
                    c.alignment = Alignment(horizontal="right")
                else:
                    c.number_format = '#,##0.00'
                    c.alignment = Alignment(horizontal="right")
            elif isinstance(cell_value, int) and not isinstance(cell_value, bool):
                c.value = cell_value
                c.number_format = '#,##0'
                c.alignment = Alignment(horizontal="right")
            elif isinstance(cell_value, (datetime, date)):
                c.value = cell_value.strftime('%Y-%m-%d')
                c.alignment = Alignment(horizontal="center")
            else:
                sanitized_val = sanitize_for_spreadsheet(str(cell_value))
                c.value = sanitized_val
                # Align codes and badges to center, text to left
                if col_type == 'center':
                    c.alignment = Alignment(horizontal="center")
                else:
                    c.alignment = Alignment(horizontal="left")

        current_row += 1

    last_data_row = current_row - 1

    # 5. Freeze Panes & Auto-filter
    ws.freeze_panes = f"A{header_row_idx + 1}"
    if rows and headers:
        last_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A{header_row_idx}:{last_col_letter}{last_data_row}"

    # 6. Auto-adjust Column Widths
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(headers[col_idx - 1]))
        # Inspect sample of data cells
        for r_idx in range(header_row_idx + 1, min(last_data_row + 1, header_row_idx + 100)):
            val = ws.cell(row=r_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    # 7. Write to response
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def export_multi_sheet_xlsx(filename, title, metadata, sheets_data):
    """
    Builds a consolidated multi-sheet executive workbook.
    sheets_data is a list of dicts:
    [
        {
            'sheet_name': 'Summary',
            'headers': [...],
            'rows': [...],
            'column_types': [...]
        },
        ...
    ]
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=13, bold=True, color="1B365D")
    meta_label_font = Font(name="Calibri", size=10, bold=True, color="333333")
    meta_val_font = Font(name="Calibri", size=10, color="555555")
    cell_font = Font(name="Calibri", size=10)

    thin_border_side = Side(border_style="thin", color="D3D3D3")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(
        left=Side(border_style="thin", color="0D1E36"),
        right=Side(border_style="thin", color="0D1E36"),
        top=Side(border_style="medium", color="0D1E36"),
        bottom=Side(border_style="medium", color="0D1E36")
    )

    for item in sheets_data:
        sheet_name = item['sheet_name'][:31]
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        headers = item.get('headers', [])
        rows = item.get('rows', [])
        column_types = item.get('column_types', [])

        current_row = 1
        ws.cell(row=current_row, column=1, value=f"{title} — {sheet_name}").font = title_font
        current_row += 1
        ws.cell(row=current_row, column=1, value=f"Generated: {timezone.now():%Y-%m-%d %H:%M}").font = meta_val_font
        current_row += 2

        header_row_idx = current_row
        for col_idx, header in enumerate(headers, start=1):
            c = ws.cell(row=header_row_idx, column=col_idx, value=header)
            c.fill = navy_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = header_border

        ws.row_dimensions[header_row_idx].height = 26

        current_row += 1
        for row in rows:
            ws.row_dimensions[current_row].height = 20
            for col_idx, cell_value in enumerate(row, start=1):
                c = ws.cell(row=current_row, column=col_idx)
                c.font = cell_font
                c.border = cell_border

                col_type = column_types[col_idx - 1] if col_idx - 1 < len(column_types) else None

                if cell_value is None:
                    c.value = "—"
                    c.alignment = Alignment(horizontal="center")
                elif isinstance(cell_value, (Decimal, float)):
                    c.value = float(cell_value)
                    if col_type == 'currency':
                        c.number_format = '₱#,##0.00'
                        c.alignment = Alignment(horizontal="right")
                    else:
                        c.number_format = '#,##0.00'
                        c.alignment = Alignment(horizontal="right")
                elif isinstance(cell_value, int) and not isinstance(cell_value, bool):
                    c.value = cell_value
                    c.number_format = '#,##0'
                    c.alignment = Alignment(horizontal="right")
                elif isinstance(cell_value, (datetime, date)):
                    c.value = cell_value.strftime('%Y-%m-%d')
                    c.alignment = Alignment(horizontal="center")
                else:
                    c.value = sanitize_for_spreadsheet(str(cell_value))
                    if col_type == 'center':
                        c.alignment = Alignment(horizontal="center")
                    else:
                        c.alignment = Alignment(horizontal="left")

            current_row += 1

        last_data_row = current_row - 1
        ws.freeze_panes = f"A{header_row_idx + 1}"
        if rows and headers:
            last_col_letter = get_column_letter(len(headers))
            ws.auto_filter.ref = f"A{header_row_idx}:{last_col_letter}{last_data_row}"

        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(headers[col_idx - 1]))
            for r_idx in range(header_row_idx + 1, min(last_data_row + 1, header_row_idx + 50)):
                val = ws.cell(row=r_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
